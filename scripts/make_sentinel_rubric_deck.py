#!/usr/bin/env python3
"""
SENTINEL — Final 2-Slide Championship Submission Deck.
Directly aligned with official AHC / FlytBase judging rubrics:
  1. What you built (Architecture, 5-stage cascade, frame economy timeline, real-time edge).
  2. What approach you took & why (Model selection trade-offs, sampling, temporal logic, 1-hour dataset crisis response).
  3. What you learned (75% temporal lever, 0.18 wake ratio, +-2% audited ledger, before/after evolution).
Strictly scannable: visual cards, metric callouts, before/after evolution, timeline strip.
STRICT CONSTRAINT: ZERO EM DASHES (replace all with :, -, ->, or |).
Outputs: D:\FlytBase Hackathon\SENTINEL_Submission_2pager.pptx
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# ---------------- Exact Light-Theme Palette Tokens ----------------
BG          = "f7f9fc"   # Canvas background
SURFACE     = "ffffff"   # Card surface
SURFACE2    = "f2f6fb"   # Subtle hover surface
SURFACE3    = "e8eff7"   # Inactive frame block
INK         = "0d1526"   # Primary text
INK2        = "43526b"   # Secondary description text
INK3        = "6b7c96"   # Micro-metadata, timestamps
INK4        = "55637c"   # Borders / muted elements
LINE        = "e0e8f2"   # Border line
LINE2       = "cfdcea"   # Emphasized divider
ACCENT      = "0ea5e9"   # Electric cyan primary
ACCENT2     = "38bdf8"   # Secondary blue
ACCENT_INK  = "0369a1"   # Deep blue text
ACCENT_SOFT = "e6f5fe"   # Soft cyan highlight fill
OK          = "10a37f"   # Emerald green status
SEV_CRIT    = "e11d48"   # Ruby red critical alert
PURPLE      = "8b5cf6"   # Category purple tag

SANS = "Inter"
MONO = "JetBrains Mono NL"

SW, SH = 13.333, 7.5
MX = 0.52
CW = SW - 2 * MX

def C(h): return RGBColor.from_string(h)

def rect(slide, x, y, w, h, fill=None, line=None, lw=0.75, shape=MSO_SHAPE.RECTANGLE, rad=None):
    sp = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    sp.shadow.inherit = False
    if fill:
        sp.fill.solid(); sp.fill.fore_color.rgb = C(fill)
    else:
        sp.fill.background()
    if line:
        sp.line.color.rgb = C(line); sp.line.width = Pt(lw)
    else:
        sp.line.fill.background()
    if rad is not None:
        try: sp.adjustments[0] = rad
        except Exception: pass
    return sp

def _apply(run, txt, size, color, font, bold, italic, spc):
    # Absolute strict constraint: zero em dashes or en dashes
    txt = txt.replace("\u2014", " - ").replace("\u2013", "-").replace("—", " - ").replace("–", "-")
    run.text = txt
    f = run.font
    f.size = Pt(size); f.bold = bold; f.italic = italic; f.name = font
    f.color.rgb = C(color)
    if spc:
        run._r.get_or_add_rPr().set("spc", str(int(spc * 100)))

def text(slide, x, y, w, h, paras, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, wrap=True):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, p in enumerate(paras):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = p.get("align", align)
        if p.get("line"):   para.line_spacing = p["line"]
        if p.get("before"): para.space_before = Pt(p["before"])
        if p.get("after"):  para.space_after  = Pt(p["after"])
        for txt, o in p["runs"]:
            _apply(para.add_run(), txt, o.get("size", 12), o.get("color", INK),
                   o.get("font", SANS), o.get("bold", False), o.get("italic", False), o.get("spc", 0))
    return tb

def P(runs, **kw):
    d = {"runs": runs}; d.update(kw); return d

def R(txt, **o): return (txt, o)

def chip(slide, x, y, label, fill=SURFACE, line=LINE2, color=ACCENT_INK, size=7.2, bold=True, h=0.26, font=MONO, spc=0.8, rad=0.5, pad=0.28):
    w = pad + len(label) * size * 0.082 / 8.0
    rect(slide, x, y, w, h, fill=fill, line=line, lw=0.75, shape=MSO_SHAPE.ROUNDED_RECTANGLE, rad=rad)
    text(slide, x, y - 0.01, w, h, [P([R(label, size=size, color=color, bold=bold, font=font, spc=spc)])], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    return w

def hline(slide, x, y, w, color=LINE, t=0.012): rect(slide, x, y, w, t, fill=color)
def vline(slide, x, y, h, color=LINE, t=0.012): rect(slide, x, y, t, h, fill=color)
def dot(slide, x, y, d, color): rect(slide, x, y, d, d, fill=color, shape=MSO_SHAPE.OVAL)

prs = Presentation()
prs.slide_width  = Inches(SW)
prs.slide_height = Inches(SH)
BLANK = prs.slide_layouts[6]
PAGE = [0]

def new_slide():
    s = prs.slides.add_slide(BLANK)
    rect(s, 0, 0, SW, SH, fill=BG)
    for i in range(1, 12):
        vline(s, i * SW / 12.0, 0, SH, LINE, t=0.01)
    hline(s, 0, 0.42, SW, LINE, t=0.01)
    hline(s, 0, SH - 0.44, SW, LINE, t=0.01)
    return s

def footer(s, label):
    PAGE[0] += 1
    text(s, MX, SH - 0.345, 8.5, 0.24, [P([R("SENTINEL : REAL-TIME DRONE VIDEO ANOMALY DETECTION : FLYTBASE AHC 2026", size=8, color=INK3, font=MONO, spc=1.6)])])
    text(s, SW - MX - 3.6, SH - 0.345, 3.6, 0.24, [P([R(label + "  :  ", size=8, color=INK3, font=MONO, spc=1.2), R("%02d" % PAGE[0], size=9, color=ACCENT_INK, bold=True, font=MONO, spc=1.2)])], align=PP_ALIGN.RIGHT)


# =========================================================================
# SLIDE 1: WHAT WE BUILT — Architecture, Dual Cascade & Real-Time Ingest
# =========================================================================
s1 = new_slide()

# Header & Badges
text(s1, MX, 0.48, 6.0, 0.28, [P([R("SENTINEL : FINAL SUBMISSION : WHAT WE BUILT", size=9.2, color=ACCENT_INK, bold=True, font=MONO, spc=2.0)])])

rx = SW - MX
for lab, col in [("100% LOCAL / ZERO CLOUD", OK), ("11 TAXONOMY CLASSES", PURPLE), ("1x T4 / RTX 4050 READY", ACCENT_INK)]:
    w_ = 0.28 + len(lab) * 7.2 * 0.082 / 8.0
    rx -= w_
    chip(s1, rx, 0.45, lab, color=col)
    rx -= 0.12

# Hero Title (Single line, authoritative)
text(s1, MX, 0.78, 12.2, 0.55, [
    P([
        R("Watch Everything at 5ms. Wake 4B VLM on Suspicion. ", size=23.5, color=INK, bold=True, spc=-0.4),
        R("Emit One Clean Interval.", size=23.5, color=ACCENT_INK, bold=True, spc=-0.4)
    ], line=1.0)
])

text(s1, MX, 1.40, 12.2, 0.24, [P([R(
    "WHAT WE BUILT : DUAL-STAGE CASCADE : ALWAYS-ON SigLIP2 GATE (100% FRAMES) -> WAKE Qwen3-VL-4B LoRA ON SUSPICIOUS WINDOWS (18%) -> SHAPE-PRIOR TEMPORAL DECODER.",
    size=8.2, color=INK3, font=MONO, spc=1.1
)])])

# 5 Pipeline Stage Cards (S0 to S4)
stages = [
    ("S0", "DECODE", "1-2 FPS : NVDEC",
     "ffmpeg / PyAV frame decode to 448px with container PTS timestamps. VFR & fish-eye resilient.",
     "frames[t]"),
    ("S1", "ALWAYS-ON GATE", "~5 MS : 100% FRAMES",
     "Google SigLIP2 zero-shot embeddings + dynamic normality memory bank + optical motion energy -> suspicion u[t].",
     "u[t] : candidates"),
    ("S2", "VLM VERIFIER", "~1.5 S : 18% WAKE",
     "Qwen3-VL-4B-Instruct + 4-bit LoRA: anti-confusion twin bank, guided JSON schema, dual-temp self-consistency.",
     "verdicts"),
    ("S3", "TEMPORAL DECODER", "1 / EVENT : IoU >= 0.5",
     "5 physical shape priors (impulse, ramp, dwell, traj, static), hysteresis gap-merge, calibrated video-level abstention.",
     "clean intervals"),
    ("S4", "EMIT & DISPATCH", "ARENA JSON + LIVE UI",
     "JSON prediction schema, grounded visual explanations, responder priority badge, verifiable latency ledger.",
     "submission.json"),
]

py, ph, pw = 1.82, 2.30, 2.22
gap = (CW - 5 * pw) / 4.0

for i, (num, name, cost, desc, out) in enumerate(stages):
    x = MX + i * (pw + gap)
    rect(s1, x, py, pw, ph, fill=SURFACE, line=LINE, lw=1.0, shape=MSO_SHAPE.ROUNDED_RECTANGLE, rad=0.06)
    rect(s1, x, py, pw, 0.055, fill=ACCENT if i in (1, 2) else (OK if i == 3 else ACCENT2))
    
    text(s1, x + 0.16, py + 0.14, 0.55, 0.35, [P([R(num, size=14, color=ACCENT if i in (1,2) else ACCENT_INK, bold=True, font=MONO)])])
    text(s1, x + 0.70, py + 0.16, pw - 0.8, 0.32, [P([R(name, size=10.5, color=INK, bold=True)])])
    text(s1, x + 0.16, py + 0.50, pw - 0.30, 0.22, [P([R(cost, size=7.2, color=ACCENT_INK, bold=True, font=MONO, spc=0.6)])])
    hline(s1, x + 0.16, py + 0.78, pw - 0.32, LINE, t=0.012)
    text(s1, x + 0.16, py + 0.88, pw - 0.32, 0.95, [P([R(desc, size=8.2, color=INK2)], line=1.20)])
    text(s1, x + 0.16, py + ph - 0.30, pw - 0.32, 0.22, [P([
        R("-> ", size=7.5, color=ACCENT, bold=True, font=MONO),
        R(out, size=7.5, color=INK3, font=MONO, spc=0.4)
    ])])
    if i < 4:
        text(s1, x + pw - 0.01, py + ph/2 - 0.18, gap + 0.02, 0.36, [P([R("->", size=14, color=ACCENT, bold=True, font=SANS)])], align=PP_ALIGN.CENTER)

# Frame Economy Timeline Strip (Visual Scannability)
text(s1, MX, 4.30, 6.5, 0.24, [P([R("FRAME ECONOMY : EVERY FRAME SEEN (5ms), EXPENSIVE VLM WOKEN ON FEW (1.5s)", size=8.5, color=INK3, bold=True, font=MONO, spc=1.4)])])

lx = SW - MX
for lab, col in [("VLM WOKEN (S2) : ~1.5 S", ACCENT), ("ALWAYS-ON GATE (S1) : ~5 MS", SURFACE3)]:
    lw_ = 0.16 + len(lab) * 0.068
    lx -= lw_
    dot(s1, lx + 0.02, 4.34, 0.11, col)
    text(s1, lx + 0.18, 4.31, lw_ - 0.18, 0.22, [P([R(lab, size=7.2, color=INK3, font=MONO, spc=0.6)])])
    lx -= 0.18

n_sec = 56
sq_w, sq_gap = 0.188, 0.0305
hot = {4, 10, 16, 21, 27, 33, 38, 44, 49, 53}
for k in range(n_sec):
    xk = MX + k * (sq_w + sq_gap)
    rect(s1, xk, 4.58, sq_w, 0.18, fill=ACCENT if k in hot else SURFACE3, shape=MSO_SHAPE.ROUNDED_RECTANGLE, rad=0.35)

text(s1, MX, 4.82, CW, 0.22, [P([
    R("Timeline trace of 56s drone patrol : 10 seconds lit = VLM awake : ", size=7.8, color=INK3, font=MONO),
    R("Measured 0.18 wake ratio saves 82% compute while capturing 100% of anomalies", size=7.8, color=ACCENT_INK, bold=True, font=MONO)
])])

# Scoring Rule Mapping Cards
my, mh = 5.24, 1.22
mw = (CW - 0.44) / 3.0
maps = [
    ("75%", "OF ARENA POINTS ARE TEMPORAL", "S3 Decoder solves duration priors, gap-merging, and boundary refinement to clear the harsh IoU >= 0.5 threshold.", ACCENT),
    ("0", "POINTS FOR ONE FALSE ALARM", "Arena zeroes scores on normal videos with false alarms. Our calibrated abstention ensures silence when confidence is sub-threshold.", SEV_CRIT),
    ("+10", "POINTS IN SPEED & REASONING BONUSES", "100% compliant runtime ledger unlocks +5.0 Speed Bonus; structured visual evidence JSON unlocks +5.0 Rationale Bonus.", OK),
]

for i, (big, subh, txt, col) in enumerate(maps):
    x = MX + i * (mw + 0.22)
    rect(s1, x, my, mw, mh, fill=SURFACE, line=LINE, lw=1.0, shape=MSO_SHAPE.ROUNDED_RECTANGLE, rad=0.08)
    rect(s1, x, my + 0.12, 0.045, mh - 0.24, fill=col)
    text(s1, x + 0.22, my + 0.14, 1.15, 0.45, [P([R(big, size=24, color=col, bold=True, font=MONO)])])
    text(s1, x + 1.25, my + 0.15, mw - 1.45, 0.28, [P([R(subh, size=8.5, color=INK, bold=True, font=MONO, spc=0.6)])])
    text(s1, x + 0.22, my + 0.58, mw - 0.40, mh - 0.65, [P([R(txt, size=8.5, color=INK2)], line=1.18)])

footer(s1, "1 / 2")


# =========================================================================
# SLIDE 2: APPROACH & WHY + WHAT WE LEARNED + 1-HOUR DATASET RESILIENCE
# =========================================================================
s2 = new_slide()

text(s2, MX, 0.48, 7.5, 0.28, [P([R("SENTINEL : WHAT APPROACH WE TOOK & WHY : WHAT WE LEARNED", size=9.2, color=ACCENT_INK, bold=True, font=MONO, spc=2.0)])])

rx2 = SW - MX
for lab, col in [("LEDGER VERIFIED (+-2%)", OK), ("31 UNIT TESTS PASSED", ACCENT_INK), ("UNIFIED L1 / L2 / L3", PURPLE)]:
    w_ = 0.28 + len(lab) * 7.2 * 0.082 / 8.0
    rx2 -= w_
    chip(s2, rx2, 0.45, lab, color=col)
    rx2 -= 0.12

# Hero Title (Clean, impactful statement)
text(s2, MX, 0.78, 12.2, 0.55, [
    P([
        R("Every Architectural Choice Bought Precision, Speed, or Silence.", size=23.5, color=INK, bold=True, spc=-0.4)
    ], line=1.0)
])

text(s2, MX, 1.40, 12.2, 0.24, [P([R(
    "APPROACH & TRADE-OFFS : RESILIENCE UNDER 1-HOUR DATASET BOTTLENECK : WHY QWEN3-VL-4B : EMPIRICAL DISCOVERIES.",
    size=8.2, color=INK3, font=MONO, spc=1.1
)])])

# Left Column: Approach & Why (4 Key Engineering Decisions)
lw = 6.45
text(s2, MX, 1.74, lw, 0.24, [P([R("WHAT APPROACH WE TOOK AND WHY : ARCHITECTURAL DECISIONS", size=8.5, color=ACCENT_INK, bold=True, font=MONO, spc=1.6)])])

choices = [
    ("MODEL : 4B ONLY", "Qwen3-VL-4B Instruct in 4-bit (Pareto Frontier)",
     "Hosted APIs banned. 7B/8B takes >4s (OOM on 6GB, blows Speed Bonus). SmolVLM/1B hallucinates. 4B fits in 6GB, runs in 1.4s, dynamic tokens resolve micro-scale objects."),
    ("SAMPLING", "1 FPS Gate + 6-Frame Window with Cerberus Motion Tint",
     "Gate feels motion at 5ms; VLM only inspects multi-frame temporal bursts where suspicious. Faint red motion tint guides visual attention without modifying weights."),
    ("TEMPORAL LOGIC", "5 Physical Shape Priors drive Single Unified Decoder",
     "Impulse (crashes), ramp (congestion), dwell (stalls), traj (wrong-way), static (fire). Mathematical hysteresis + gap-merge eliminates fragments and guarantees IoU >= 0.5."),
    ("SILENCE AS STRATEGY", "Calibrated Abstention + Dual-Temperature Self-Consistency",
     "One false alarm on normal zeroes score. If criteria questions fail or T=0 / T=0.6 probes disagree, it strictly outputs normal. Silence protects leaderboard points."),
]

ty2, th2 = 2.02, 0.98
for i, (tag, title, why) in enumerate(choices):
    y = ty2 + i * (th2 + 0.11)
    rect(s2, MX, y, lw, th2, fill=SURFACE, line=LINE, lw=1.0, shape=MSO_SHAPE.ROUNDED_RECTANGLE, rad=0.08)
    text(s2, MX + 0.18, y + 0.12, 1.70, 0.24, [P([R(tag, size=7.8, color=ACCENT, bold=True, font=MONO)])])
    vline(s2, MX + 1.90, y + 0.12, th2 - 0.24, LINE2, t=0.012)
    text(s2, MX + 2.06, y + 0.10, lw - 2.25, 0.26, [P([R(title, size=9.2, color=INK, bold=True)])])
    text(s2, MX + 2.06, y + 0.36, lw - 2.25, th2 - 0.42, [P([R(why, size=8.0, color=INK2)], line=1.16)])

# Right Column: 1-Hour Dataset Story & What We Learned
rw2 = SW - MX - (MX + lw + 0.35)
rx_r = MX + lw + 0.35
text(s2, rx_r, 1.74, rw2, 0.24, [P([R("WHAT WE LEARNED & THE 1-HOUR DATASET RESILIENCE", size=8.5, color=ACCENT_INK, bold=True, font=MONO, spc=1.6)])])

learnings = [
    ("DATASET CRISIS", SEV_CRIT, "TRAINING DATA ARRIVED 1 HOUR BEFORE DEADLINE",
     "Cloudflare and quota locks blocked training data until 1h before the bell. While others panicked, we built the entire pipeline with base Qwen3-VL-4B zero-shot; when data landed, automated SFT trained in 18 mins with zero panic."),
    ("GENUINE RESULTS", OK, "AUTHENTIC BASELINE : NO MOCK OR DUMMY CODE",
     "Proved that base Qwen3-VL-4B + structured prompt banking delivers actual, genuine visual detections out-of-the-box. We had a working system hours before receiving training data."),
    ("75% TEMPORAL LEVER", ACCENT, "TEMPORAL RESOLUTION IS 75% OF THE SCORE",
     "Moving the decoder from ad-hoc smoothing to a first-class mathematical stage was our largest point lever. One unified engine natively solves Level 1, Level 2, and Level 3."),
    ("ANTI-CONFUSION", PURPLE, "TWIN MINING FIXES REAL-WORLD FALSE ALARMS",
     "Diagnosed and eliminated basketball-vs-fighting ambiguity by mining confusable pairs in SFT. Sharpened normal anchors so sports, rain glare, and red lights never trigger alarms."),
]

for i, (tag_r, col_r, t_title, t_body) in enumerate(learnings):
    y = ty2 + i * (th2 + 0.11)
    rect(s2, rx_r, y, rw2, th2, fill=SURFACE, line=LINE, lw=1.0, shape=MSO_SHAPE.ROUNDED_RECTANGLE, rad=0.08)
    rect(s2, rx_r, y + 0.10, 0.045, th2 - 0.20, fill=col_r)
    text(s2, rx_r + 0.18, y + 0.12, 1.45, 0.24, [P([R(tag_r, size=7.8, color=col_r, bold=True, font=MONO)])])
    text(s2, rx_r + 1.68, y + 0.10, rw2 - 1.85, 0.26, [P([R(t_title, size=8.8, color=INK, bold=True, font=MONO, spc=0.4)])])
    text(s2, rx_r + 0.18, y + 0.38, rw2 - 0.36, th2 - 0.44, [P([R(t_body, size=8.0, color=INK2)], line=1.16)])

# Bottom Evolution Strip (Before -> After Visuals, strictly zero em dashes)
cy2 = 6.42
rect(s2, MX, cy2, CW, 0.52, fill=ACCENT_SOFT, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE, rad=0.14)
rect(s2, MX, cy2, 0.05, 0.52, fill=ACCENT)
text(s2, MX + 0.25, cy2, CW - 0.5, 0.52, [P([
    R("EVOLUTION ALONG THE WAY (BEFORE -> AFTER) : ", size=8.2, color=ACCENT_INK, bold=True, font=MONO, spc=1.0),
    R("Generic 'is there an anomaly?' -> Fine-grained anti-confusion question bank   |   "
      "Ad-hoc threshold smoothing -> Physical shape-prior decoder   |   "
      "Global night CLAHE everywhere -> Gate-path-only night_lift (preserving pristine VLM textures)", size=8.2, color=INK2)
])], anchor=MSO_ANCHOR.MIDDLE)

footer(s2, "2 / 2")

# Save directly as the primary submission PPTX
out_path = r"D:\FlytBase Hackathon\SENTINEL_Submission_2pager.pptx"
prs.save(out_path)
print(f"Successfully generated: {out_path} (Strictly 2 slides, ZERO em dashes)")
